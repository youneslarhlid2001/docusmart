from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.api.deps import get_db
from app.models.supplier import Supplier
from app.models.document import Document
from app.schemas.fraud import SupplierRead
from app.schemas.document import DocumentListRead

router = APIRouter()


@router.get("/suppliers", response_model=list[SupplierRead])
async def list_suppliers(
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    """Liste tous les fournisseurs connus, triés par score de conformité."""
    result = await db.execute(
        select(Supplier)
        .order_by(Supplier.compliance_score.asc())  # Les moins conformes en premier
        .offset(skip)
        .limit(limit)
    )
    return result.scalars().all()


@router.get("/suppliers/export/excel")
async def export_suppliers_excel(db: AsyncSession = Depends(get_db)):
    """Exporte le scoring de tous les fournisseurs au format Excel."""
    result = await db.execute(
        select(Supplier).order_by(Supplier.compliance_score.asc())
    )
    suppliers = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scoring fournisseurs"

    header_fill = PatternFill("solid", fgColor="6366F1")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    alt_fill = PatternFill("solid", fgColor="F8FAFC")
    green_font = Font(color="059669", bold=True)
    red_font = Font(color="DC2626", bold=True)
    orange_font = Font(color="D97706", bold=True)

    headers = ["SIREN", "SIRET", "Raison sociale", "Adresse", "Score conformité (%)", "Nb documents", "Dernier document"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, supplier in enumerate(suppliers, 2):
        fill = alt_fill if row_idx % 2 == 0 else None
        values = [
            supplier.siren,
            supplier.siret or "—",
            supplier.raison_sociale or "—",
            supplier.adresse or "—",
            round(supplier.compliance_score, 1),
            supplier.document_count,
            str(supplier.last_document_at)[:10] if supplier.last_document_at else "—",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            if fill:
                cell.fill = fill
            if col == 5:  # Score colonne
                score = supplier.compliance_score
                cell.font = green_font if score >= 80 else (orange_font if score >= 50 else red_font)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="scoring_fournisseurs.xlsx"'},
    )


@router.get("/suppliers/{siren}", response_model=SupplierRead)
async def get_supplier(siren: str, db: AsyncSession = Depends(get_db)):
    """Détail d'un fournisseur par son SIREN."""
    result = await db.execute(select(Supplier).where(Supplier.siren == siren))
    supplier = result.scalar_one_or_none()
    if not supplier:
        raise HTTPException(status_code=404, detail="Fournisseur non trouvé")
    return supplier


@router.get("/suppliers/{siren}/documents", response_model=list[DocumentListRead])
async def get_supplier_documents(
    siren: str,
    db: AsyncSession = Depends(get_db),
):
    """Liste tous les documents d'un fournisseur donné."""
    result = await db.execute(
        select(Document)
        .where(Document.supplier_siren == siren)
        .order_by(Document.created_at.desc())
    )
    return result.scalars().all()
