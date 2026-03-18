import uuid
from datetime import datetime, date, timedelta
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import select, func, update
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.api.deps import get_db
from app.models.alert import ComplianceAlert
from app.models.document import Document
from app.models.supplier import Supplier
from app.schemas.fraud import (
    AlertRead, AlertUpdateRequest, StatsRead,
    AnalyticsRead, ComplianceTrendPoint, HeatmapCell,
    SupplierRadarData, SankeyData, SankeyNode, SankeyLink, RecurringAlertSupplier,
)

router = APIRouter()


@router.get("/compliance/alerts", response_model=list[AlertRead])
async def list_alerts(
    status: str | None = None,
    severity: str | None = None,
    document_id: str | None = None,
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    """Liste toutes les incohérences détectées, filtrables par statut, sévérité et document."""
    query = select(ComplianceAlert).order_by(ComplianceAlert.created_at.desc())

    if status:
        query = query.where(ComplianceAlert.status == status)
    if severity:
        query = query.where(ComplianceAlert.severity == severity)
    if document_id:
        try:
            query = query.where(ComplianceAlert.document_id == uuid.UUID(document_id))
        except ValueError:
            pass

    query = query.offset(skip).limit(limit)
    result = await db.execute(query)
    return result.scalars().all()


@router.patch("/compliance/alerts/{alert_id}", response_model=AlertRead)
async def update_alert(
    alert_id: uuid.UUID,
    body: AlertUpdateRequest,
    db: AsyncSession = Depends(get_db),
):
    """
    Met à jour le statut d'une alerte : RESOLVED, FALSE_POSITIVE.
    Utilisé par les gestionnaires pour traiter les incohérences.
    """
    result = await db.execute(
        select(ComplianceAlert).where(ComplianceAlert.id == alert_id)
    )
    alert = result.scalar_one_or_none()
    if not alert:
        raise HTTPException(status_code=404, detail="Alerte non trouvée")

    allowed_statuses = {"RESOLVED", "FALSE_POSITIVE", "OPEN"}
    if body.status not in allowed_statuses:
        raise HTTPException(status_code=400, detail=f"Statut invalide. Valeurs : {allowed_statuses}")

    alert.status = body.status
    if body.status == "RESOLVED":
        alert.resolved_at = datetime.utcnow()
    await db.commit()
    await db.refresh(alert)
    return alert


@router.get("/compliance/export/excel")
async def export_alerts_excel(db: AsyncSession = Depends(get_db)):
    """Exporte toutes les alertes de conformité au format Excel."""
    result = await db.execute(
        select(ComplianceAlert).order_by(ComplianceAlert.created_at.desc())
    )
    alerts = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alertes de conformité"

    header_fill = PatternFill("solid", fgColor="6366F1")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    alt_fill = PatternFill("solid", fgColor="F8FAFC")

    headers = ["ID", "Type alerte", "Sévérité", "Statut", "Description", "Document ID", "Créé le", "Résolu le"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, alert in enumerate(alerts, 2):
        fill = alt_fill if row_idx % 2 == 0 else None
        values = [
            str(alert.id),
            alert.alert_type or "—",
            alert.severity,
            alert.status,
            alert.description or "—",
            str(alert.document_id),
            str(alert.created_at)[:19] if alert.created_at else "—",
            str(alert.resolved_at)[:19] if alert.resolved_at else "—",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            if fill:
                cell.fill = fill

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="alertes_conformite.xlsx"'},
    )


@router.get("/stats", response_model=StatsRead)
async def get_stats(db: AsyncSession = Depends(get_db)):
    """KPIs du dashboard : volumes, taux de conformité, alertes actives."""
    # Comptes documents par statut
    total_docs_result = await db.execute(select(func.count()).select_from(Document))
    total = total_docs_result.scalar_one()

    done_result = await db.execute(
        select(func.count()).select_from(Document).where(Document.status == "DONE")
    )
    done = done_result.scalar_one()

    error_result = await db.execute(
        select(func.count()).select_from(Document).where(Document.status == "ERROR")
    )
    errors = error_result.scalar_one()

    processing_statuses = ["UPLOADED", "OCR_PROCESSING", "CLASSIFYING", "EXTRACTING", "VERIFYING"]
    processing_result = await db.execute(
        select(func.count()).select_from(Document).where(Document.status.in_(processing_statuses))
    )
    processing = processing_result.scalar_one()

    # Alertes
    total_alerts_result = await db.execute(select(func.count()).select_from(ComplianceAlert))
    total_alerts = total_alerts_result.scalar_one()

    open_alerts_result = await db.execute(
        select(func.count()).select_from(ComplianceAlert).where(ComplianceAlert.status == "OPEN")
    )
    open_alerts = open_alerts_result.scalar_one()

    # Documents par type
    types_result = await db.execute(
        select(Document.file_type, func.count())
        .where(Document.file_type.isnot(None))
        .group_by(Document.file_type)
    )
    docs_by_type = {row[0]: row[1] for row in types_result.all()}

    # Taux de conformité : documents sans alerte OPEN sur total traités
    docs_with_alerts_result = await db.execute(
        select(func.count(ComplianceAlert.document_id.distinct()))
        .where(ComplianceAlert.status == "OPEN")
    )
    docs_with_open_alerts = docs_with_alerts_result.scalar_one()
    compliance_rate = ((done - docs_with_open_alerts) / done * 100) if done > 0 else 100.0

    # Volume 7 derniers jours (simplifié)
    from sqlalchemy import cast, Date, text
    from datetime import date, timedelta
    last_7_days = []
    for i in range(6, -1, -1):
        day = date.today() - timedelta(days=i)
        count_result = await db.execute(
            select(func.count()).select_from(Document).where(
                func.date(Document.created_at) == day
            )
        )
        last_7_days.append({"date": day.isoformat(), "count": count_result.scalar_one()})

    return StatsRead(
        total_documents=total,
        documents_done=done,
        documents_processing=processing,
        documents_error=errors,
        total_alerts=total_alerts,
        open_alerts=open_alerts,
        compliance_rate=round(compliance_rate, 1),
        documents_by_type=docs_by_type,
        documents_last_7_days=last_7_days,
    )


@router.get("/analytics", response_model=AnalyticsRead)
async def get_analytics(
    months: int = Query(6, ge=1, le=12),
    db: AsyncSession = Depends(get_db),
):
    """Données analytiques avancées pour le tableau de bord stratégique."""
    today = date.today()

    # ── Génère la liste des N derniers mois ─────────────────────────────
    months_list: list[str] = []
    d = today.replace(day=1)
    for _ in range(months):
        months_list.append(d.strftime("%Y-%m"))
        d = (d - timedelta(days=1)).replace(day=1)
    months_list.reverse()

    # ── 1. Compliance trend (taux conformité par mois) ───────────────────
    # Documents DONE par mois
    done_by_month_result = await db.execute(
        select(
            func.to_char(Document.created_at, "YYYY-MM").label("month"),
            func.count().label("cnt"),
        )
        .where(Document.status == "DONE")
        .group_by("month")
    )
    done_by_month = {row.month: row.cnt for row in done_by_month_result.all()}

    # Docs avec alertes OPEN par mois
    alerts_by_month_result = await db.execute(
        select(
            func.to_char(Document.created_at, "YYYY-MM").label("month"),
            func.count(ComplianceAlert.document_id.distinct()).label("cnt"),
        )
        .join(Document, ComplianceAlert.document_id == Document.id)
        .where(ComplianceAlert.status == "OPEN")
        .group_by("month")
    )
    alerts_by_month = {row.month: row.cnt for row in alerts_by_month_result.all()}

    # Alertes créées par mois (total)
    created_alerts_by_month_result = await db.execute(
        select(
            func.to_char(ComplianceAlert.created_at, "YYYY-MM").label("month"),
            func.count().label("cnt"),
        )
        .group_by("month")
    )
    created_alerts_by_month = {row.month: row.cnt for row in created_alerts_by_month_result.all()}

    compliance_trend = []
    for m in months_list:
        done = done_by_month.get(m, 0)
        with_alerts = alerts_by_month.get(m, 0)
        rate = round(((done - with_alerts) / done * 100), 1) if done > 0 else 100.0
        compliance_trend.append(ComplianceTrendPoint(
            month=m,
            compliance_rate=rate,
            documents_count=done,
            alerts_count=created_alerts_by_month.get(m, 0),
        ))

    # ── 2. Anomaly heatmap (alertes par type de doc × mois) ─────────────
    heatmap_result = await db.execute(
        select(
            func.to_char(ComplianceAlert.created_at, "YYYY-MM").label("month"),
            Document.file_type.label("doc_type"),
            func.count().label("cnt"),
        )
        .join(Document, ComplianceAlert.document_id == Document.id)
        .where(Document.file_type.isnot(None))
        .group_by("month", Document.file_type)
    )
    anomaly_heatmap = [
        HeatmapCell(month=row.month, doc_type=row.doc_type, count=row.cnt)
        for row in heatmap_result.all()
        if row.month in months_list
    ]

    # ── 3. Supplier radar (top 6 fournisseurs) ───────────────────────────
    suppliers_result = await db.execute(
        select(Supplier).order_by(Supplier.document_count.desc()).limit(6)
    )
    suppliers = suppliers_result.scalars().all()

    supplier_radar = []
    for s in suppliers:
        total_alerts_res = await db.execute(
            select(func.count())
            .select_from(ComplianceAlert)
            .join(Document, ComplianceAlert.document_id == Document.id)
            .where(Document.supplier_siren == s.siren)
        )
        open_alerts_res = await db.execute(
            select(func.count())
            .select_from(ComplianceAlert)
            .join(Document, ComplianceAlert.document_id == Document.id)
            .where(Document.supplier_siren == s.siren, ComplianceAlert.status == "OPEN")
        )
        supplier_radar.append(SupplierRadarData(
            siren=s.siren,
            name=s.raison_sociale or s.siren,
            compliance_score=s.compliance_score,
            doc_count=s.document_count,
            alert_count=total_alerts_res.scalar_one(),
            open_alert_count=open_alerts_res.scalar_one(),
        ))

    # ── 4. Document flow (Sankey: type doc → type alerte → sévérité) ────
    flow_result = await db.execute(
        select(
            Document.file_type,
            ComplianceAlert.alert_type,
            ComplianceAlert.severity,
            func.count().label("cnt"),
        )
        .join(Document, ComplianceAlert.document_id == Document.id)
        .where(Document.file_type.isnot(None), ComplianceAlert.alert_type.isnot(None))
        .group_by(Document.file_type, ComplianceAlert.alert_type, ComplianceAlert.severity)
    )
    flow_rows = flow_result.all()

    doc_types = list({r.file_type for r in flow_rows})
    alert_types = list({r.alert_type for r in flow_rows})
    severities = list({r.severity for r in flow_rows})

    sankey_nodes = [SankeyNode(name=n) for n in doc_types + alert_types + severities]
    node_index = {n: i for i, n in enumerate(doc_types + alert_types + severities)}

    sankey_links: list[SankeyLink] = []
    # doc_type → alert_type aggregation
    dt_at: dict[tuple, int] = {}
    at_sev: dict[tuple, int] = {}
    for r in flow_rows:
        dt_at[(r.file_type, r.alert_type)] = dt_at.get((r.file_type, r.alert_type), 0) + r.cnt
        at_sev[(r.alert_type, r.severity)] = at_sev.get((r.alert_type, r.severity), 0) + r.cnt

    for (dt, at), v in dt_at.items():
        if dt in node_index and at in node_index:
            sankey_links.append(SankeyLink(source=node_index[dt], target=node_index[at], value=v))
    for (at, sev), v in at_sev.items():
        if at in node_index and sev in node_index:
            sankey_links.append(SankeyLink(source=node_index[at], target=node_index[sev], value=v))

    document_flow = SankeyData(nodes=sankey_nodes, links=sankey_links)

    # ── 5. Recurring alerts (fournisseurs problématiques) ────────────────
    recurring_result = await db.execute(
        select(
            Document.supplier_siren,
            Supplier.raison_sociale,
            func.count(ComplianceAlert.id).label("total"),
            func.sum(
                func.cast(ComplianceAlert.status == "OPEN", func.Integer)
            ).label("open_count"),
            func.max(ComplianceAlert.created_at).label("last_alert"),
        )
        .join(Document, ComplianceAlert.document_id == Document.id)
        .outerjoin(Supplier, Document.supplier_siren == Supplier.siren)
        .where(Document.supplier_siren.isnot(None))
        .group_by(Document.supplier_siren, Supplier.raison_sociale)
        .order_by(func.count(ComplianceAlert.id).desc())
        .limit(5)
    )

    recurring_alerts = []
    for row in recurring_result.all():
        recurring_alerts.append(RecurringAlertSupplier(
            siren=row.supplier_siren,
            name=row.raison_sociale or row.supplier_siren,
            total_alerts=row.total,
            open_alerts=int(row.open_count or 0),
            last_alert_at=row.last_alert,
        ))

    return AnalyticsRead(
        compliance_trend=compliance_trend,
        anomaly_heatmap=anomaly_heatmap,
        supplier_radar=supplier_radar,
        document_flow=document_flow,
        recurring_alerts=recurring_alerts,
    )
